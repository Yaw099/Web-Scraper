from pathlib import Path, PurePosixPath
from urllib.parse import (
    parse_qsl,
    quote,
    unquote,
    urlencode,
    urljoin,
    urlparse,
    urlsplit,
    urlunsplit,
)
from collections.abc import Callable
from hashlib import sha256
import requests
import pandas as pd
from pypdf import PdfReader
from docx import Document
import shutil
import subprocess
from bs4 import BeautifulSoup
from pptx import Presentation
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree as ET
from posixpath import normpath
from pptx.enum.shapes import MSO_SHAPE_TYPE

DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".ppt",
    ".pptx",
    ".txt",
    ".xls",
    ".xlsx",
    ".csv",
    ".rtf",
}

ARCHIVE_EXTENSIONS = {
    ".zip",
}

MAX_ARCHIVE_FILES = 1000

MAX_ARCHIVE_UNCOMPRESSED_SIZE = (
    1024 * 1024 * 1024
)  # 1 GB

def discover_document_links(html, base_url):
    soup = BeautifulSoup(html, "lxml")
    document_urls = []
    seen_urls = set()

    for tag in soup.find_all("a", href=True):
        absolute_url = urljoin(base_url, tag["href"])

        if not is_downloadable_url(absolute_url):
            continue

        url_key = normalize_url_for_dedup(absolute_url)

        if url_key in seen_urls:
            continue

        seen_urls.add(url_key)
        document_urls.append(absolute_url)

    return document_urls


def normalize_url_for_dedup(url: str) -> str:
    """Return a stable comparison key without changing the requested URL."""

    parts = urlsplit(str(url).strip())
    normalized_path = quote(
        unquote(parts.path),
        safe="/:@!$&'()*+,;=-._~",
    )
    normalized_query = urlencode(
        sorted(parse_qsl(parts.query, keep_blank_values=True)),
        doseq=True,
    )

    return urlunsplit((
        parts.scheme.lower(),
        parts.netloc.lower(),
        normalized_path,
        normalized_query,
        "",
    ))


def file_sha256(path, block_size=1024 * 1024):
    """Hash a file incrementally so identical downloads can be reused."""

    digest = sha256()

    with Path(path).open("rb") as file:
        for block in iter(lambda: file.read(block_size), b""):
            digest.update(block)

    return digest.hexdigest()

def find_libreoffice():
    # First check PATH
    path = shutil.which("soffice") or shutil.which("libreoffice")
    if path:
        return path

    # Common Windows install locations
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]

    for candidate in candidates:
        if Path(candidate).exists():
            return candidate

    return None

def normalize_document(path):
    path = Path(path)
    ext = path.suffix.lower()

    conversions = {
        ".doc": "docx",
        ".xls": "xlsx",
        ".ppt": "pptx",
    }

    if ext not in conversions:
        return path

    libreoffice_path = find_libreoffice()

    if not libreoffice_path:
        raise RuntimeError(
            f"Legacy Office file detected: {path.name}. "
            "LibreOffice is required to convert this file."
        )

    output_format = conversions[ext]
    output_folder = path.parent

    subprocess.run(
        [
            libreoffice_path,
            "--headless",
            "--convert-to",
            output_format,
            "--outdir",
            str(output_folder),
            str(path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )

    converted_path = path.with_suffix(f".{output_format}")

    if not converted_path.exists():
        raise FileNotFoundError(f"Converted file not found: {converted_path}")
    
    # Conversion succeeded, remove the legacy file.
    try:
        path.unlink()
    except Exception:
        # Not critical if cleanup fails.
        pass

    return converted_path

def get_url_extension(url: str) -> str:
    parsed = urlparse(url)
    decoded_path = unquote(parsed.path)

    return Path(decoded_path).suffix.lower()


def is_document_url(url: str) -> bool:
    return get_url_extension(url) in DOCUMENT_EXTENSIONS


def is_archive_url(url: str) -> bool:
    return get_url_extension(url) in ARCHIVE_EXTENSIONS


def is_downloadable_url(url: str) -> bool:
    return is_document_url(url) or is_archive_url(url)


def download_document(url, output_folder):
    output_folder = Path(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)

    parsed = urlparse(url)
    filename = Path(unquote(parsed.path)).name

    if not filename:
        filename = "downloaded_document"

    local_path = output_folder / filename

    # Skip download if we already have the original
    if local_path.exists():
        return local_path

    # Skip download if a converted Office document already exists
    suffix = local_path.suffix.lower()

    if suffix == ".doc":
        converted = local_path.with_suffix(".docx")
    elif suffix == ".xls":
        converted = local_path.with_suffix(".xlsx")
    elif suffix == ".ppt":
        converted = local_path.with_suffix(".pptx")
    else:
        converted = None

    if converted and converted.exists():
        return converted
    
    response = requests.get(
        url,
        stream=True,
        timeout=30,
        headers={
            "User-Agent": "Mozilla/5.0"
        }
    )
    response.raise_for_status()

    with open(local_path, "wb") as file:
        for chunk in response.iter_content(chunk_size=8192):
            if chunk:
                file.write(chunk)

    return local_path

def is_supported_document_file(path):
    path = Path(path)

    return (
        path.is_file()
        and path.suffix.lower() in DOCUMENT_EXTENSIONS
    )


def is_ignored_archive_member(member_path):
    parts = member_path.parts

    if "__MACOSX" in parts:
        return True

    if member_path.name.startswith("._"):
        return True

    if member_path.name in {".DS_Store", "Thumbs.db"}:
        return True

    return False


def validate_archive_member(member_name):
    member_path = PurePosixPath(member_name)

    if member_path.is_absolute():
        raise ValueError(
            f"Archive contains an absolute path: {member_name}"
        )

    if ".." in member_path.parts:
        raise ValueError(
            f"Archive contains an unsafe path: {member_name}"
        )

    return member_path


def safe_extract_zip(zip_path, output_folder):
    zip_path = Path(zip_path)
    output_folder = Path(output_folder)

    extraction_folder = output_folder / zip_path.stem
    extraction_folder.mkdir(parents=True, exist_ok=True)

    extraction_root = extraction_folder.resolve()

    extracted_files = []
    total_uncompressed_size = 0

    try:
        with ZipFile(zip_path, "r") as archive:
            members = archive.infolist()

            if len(members) > MAX_ARCHIVE_FILES:
                raise ValueError(
                    "Archive contains too many entries: "
                    f"{len(members)}"
                )

            for member in members:
                if member.is_dir():
                    continue

                member_path = validate_archive_member(
                    member.filename
                )

                if is_ignored_archive_member(member_path):
                    continue

                total_uncompressed_size += member.file_size

                if (
                    total_uncompressed_size
                    > MAX_ARCHIVE_UNCOMPRESSED_SIZE
                ):
                    raise ValueError(
                        "Archive exceeds the maximum allowed "
                        "uncompressed size."
                    )

                destination = (
                    extraction_folder
                    / Path(*member_path.parts)
                )

                resolved_destination = destination.resolve()

                if (
                    resolved_destination != extraction_root
                    and extraction_root
                    not in resolved_destination.parents
                ):
                    raise ValueError(
                        "Archive member would be extracted "
                        f"outside the target folder: "
                        f"{member.filename}"
                    )

                destination.parent.mkdir(
                    parents=True,
                    exist_ok=True,
                )

                if (
                    destination.exists()
                    and destination.stat().st_size == member.file_size
                ):
                    extracted_files.append(destination)
                    continue

                with archive.open(member, "r") as source:
                    with destination.open("wb") as target:
                        shutil.copyfileobj(source, target)

                extracted_files.append(destination)

    except BadZipFile as error:
        raise ValueError(
            f"Invalid or corrupted ZIP archive: {zip_path.name}"
        ) from error

    return extracted_files

def prepare_downloaded_files(path, output_folder):
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in DOCUMENT_EXTENSIONS:
        return [path]

    if suffix == ".zip":
        extracted_files = safe_extract_zip(
            path,
            output_folder,
        )

        supported_files = [
            extracted_file
            for extracted_file in extracted_files
            if is_supported_document_file(extracted_file)
        ]

        print(
            f"Extracted {len(extracted_files)} file(s) "
            f"from {path.name}"
        )

        print(
            f"Found {len(supported_files)} supported "
            "document file(s)"
        )

        return supported_files

    return []


def get_document_type(path):
    return Path(path).suffix.lower()


def extract_txt(path):
    path = Path(path)

    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="latin-1")


def extract_pdf(path):
    reader = PdfReader(path)

    pages_text = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text()

        if text:
            pages_text.append(f"\n--- Page {page_number} ---\n{text}")

    return "\n".join(pages_text)


def extract_docx(path):
    doc = Document(path)

    output = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            output.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        output.append(f"\n--- Table {table_index} ---")

        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = " ".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )
                cells.append(cell_text)

            # Avoid duplicate repeated cells caused by merged Word cells
            cleaned_cells = []
            for cell in cells:
                if not cleaned_cells or cleaned_cells[-1] != cell:
                    cleaned_cells.append(cell)

            output.append(" | ".join(cleaned_cells))

    return "\n".join(output)

DRAWINGML_NAMESPACE = (
    "http://schemas.openxmlformats.org/drawingml/2006/main"
)

PACKAGE_RELATIONSHIP_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)

DIAGRAM_DATA_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/"
    "officeDocument/2006/relationships/diagramData"
)

def clean_pptx_text(text):
    if not text:
        return ""

    text = text.replace("\x0b", "\n")
    text = text.replace("\r", "\n")

    cleaned_lines = []

    for line in text.splitlines():
        cleaned_line = " ".join(line.split())

        if cleaned_line:
            cleaned_lines.append(cleaned_line)

    return "\n".join(cleaned_lines)


def extract_pptx_shape_text(shape):
    output = []

    # Recursively inspect grouped shapes.
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child_shape in shape.shapes:
            output.extend(extract_pptx_shape_text(child_shape))

        return output

    if getattr(shape, "has_text_frame", False):
        text = clean_pptx_text(shape.text)

        if text:
            output.append(text)

    if getattr(shape, "has_table", False):
        table_output = ["--- Table ---"]

        for row in shape.table.rows:
            cells = [
                clean_pptx_text(cell.text).replace("\n", " ")
                for cell in row.cells
            ]

            table_output.append(" | ".join(cells))

        output.append("\n".join(table_output))

    return output

def resolve_pptx_relationship_target(target):
    normalized = normpath(f"ppt/slides/{target}")

    # normpath uses forward slashes here because this is an archive path.
    return normalized.lstrip("/")


def extract_smartart_xml_text(archive, diagram_path):
    try:
        xml_content = archive.read(diagram_path)
    except KeyError:
        return []

    root = ET.fromstring(xml_content)

    text_tag = f"{{{DRAWINGML_NAMESPACE}}}t"
    output = []

    for element in root.iter(text_tag):
        text = clean_pptx_text(element.text)

        if text:
            output.append(text)

    return output


def extract_slide_smartart_text(archive, slide_index):
    relationships_path = (
        f"ppt/slides/_rels/slide{slide_index}.xml.rels"
    )

    try:
        relationships_content = archive.read(relationships_path)
    except KeyError:
        return []

    root = ET.fromstring(relationships_content)

    relationship_tag = (
        f"{{{PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationship"
    )

    output = []

    for relationship in root.findall(relationship_tag):
        relationship_type = relationship.get("Type", "")
        target = relationship.get("Target", "")

        if relationship_type != DIAGRAM_DATA_RELATIONSHIP:
            continue

        if not target:
            continue

        diagram_path = resolve_pptx_relationship_target(target)

        output.extend(
            extract_smartart_xml_text(
                archive,
                diagram_path,
            )
        )

    return output

def remove_duplicate_pptx_text(items):
    output = []
    seen = set()

    for item in items:
        normalized = " ".join(item.lower().split())

        if not normalized:
            continue

        if normalized in seen:
            continue

        seen.add(normalized)
        output.append(item)

    return output

def extract_pptx(path):
    path = Path(path)
    presentation = Presentation(path)

    output = []

    with ZipFile(path, "r") as archive:
        for slide_index, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            output.append(f"\n--- Slide {slide_index} ---")

            slide_output = []

            # Normal text boxes, placeholders, grouped shapes, and tables.
            for shape in slide.shapes:
                slide_output.extend(
                    extract_pptx_shape_text(shape)
                )

            # SmartArt text stored in the internal PPTX XML.
            smartart_text = extract_slide_smartart_text(
                archive,
                slide_index,
            )

            if smartart_text:
                slide_output.append("--- SmartArt ---")
                slide_output.extend(smartart_text)

            slide_output = remove_duplicate_pptx_text(
                slide_output
            )

            if slide_output:
                output.extend(slide_output)
            else:
                output.append("[No extractable text found]")

    return "\n".join(output)

def extract_excel(path):
    sheets = pd.read_excel(path, sheet_name=None)

    output = []

    for sheet_name, df in sheets.items():
        output.append(f"\n--- Sheet: {sheet_name} ---\n")
        output.append(df.to_string(index=False))

    return "\n".join(output)


def extract_csv(path):
    try:
        df = pd.read_csv(path)
    except UnicodeDecodeError:
        df = pd.read_csv(path, encoding="latin-1")

    return df.to_string(index=False)


EXTRACTORS: dict[str, Callable[[Path], str]] = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".txt": extract_txt,
    ".rtf": extract_txt,
    ".xls": extract_excel,
    ".xlsx": extract_excel,
    ".csv": extract_csv,
}

def extract_text(path):
    ext = get_document_type(path)

    extractor = EXTRACTORS.get(ext)

    if extractor is None:
        raise ValueError(f"Unsupported document type: {ext}")

    return extractor(path)
    

def process_documents_from_links(
    links,
    output_folder,
):
    download_results = download_documents_from_links(
        links,
        output_folder,
    )

    results = []

    for download_result in download_results:
        if not download_result["success"]:
            results.append(download_result)
            continue

        local_path = download_result["local_path"]

        try:
            processed = process_downloaded_document_file(
                local_path
            )

            processed["url"] = download_result["url"]
            processed["archive_path"] = (
                download_result.get("archive_path", "")
            )

            results.append(processed)

        except Exception as error:
            results.append({
                "url": download_result["url"],
                "local_path": local_path,
                "normalized_path": "",
                "file_type": get_document_type(
                    local_path
                ),
                "text": "",
                "success": False,
                "error": str(error),
                "archive_path": download_result.get(
                    "archive_path",
                    "",
                ),
            })

    return results

def process_document(url, output_folder):
    try:
        local_file = download_document(url, output_folder)
        normalized_file = normalize_document(local_file)
        text = extract_text(normalized_file)

        return {
            "url": url,
            "local_path": str(local_file),
            "normalized_path": str(normalized_file),
            "file_type": get_document_type(normalized_file),
            "text": text,
            "success": True,
            "error": None
        }

    except Exception as error:
        return {
            "url": url,
            "local_path": None,
            "normalized_path": None,
            "file_type": None,
            "text": "",
            "success": False,
            "error": str(error)
        }

def process_downloaded_document_file(path):
    normalized_file = normalize_document(path)
    text = extract_text(normalized_file)

    return {
        "local_path": str(path),
        "normalized_path": str(normalized_file),
        "file_type": get_document_type(normalized_file),
        "text": text,
        "success": True,
        "error": None,
    }

def download_documents_from_links(
    links,
    output_folder,
):
    results = []
    seen_urls = set()

    for link in links:
        url = link.get("url") if isinstance(link, dict) else link

        if not url or not is_downloadable_url(url):
            continue

        url_key = normalize_url_for_dedup(url)

        if url_key in seen_urls:
            continue

        seen_urls.add(url_key)

        try:
            downloaded_file = download_document(
                url,
                output_folder,
            )

            prepared_files = prepare_downloaded_files(
                downloaded_file,
                output_folder,
            )

            if not prepared_files:
                results.append({
                    "url": url,
                    "local_path": str(downloaded_file),
                    "normalized_path": "",
                    "file_type": get_document_type(
                        downloaded_file
                    ),
                    "text": "",
                    "success": False,
                    "error": (
                        "No supported documents were found "
                        "in the downloaded file."
                    ),
                    "archive_path": (
                        str(downloaded_file)
                        if is_archive_url(url)
                        else ""
                    ),
                })

                continue

            for prepared_file in prepared_files:
                results.append({
                    "url": url,
                    "local_path": str(prepared_file),
                    "normalized_path": "",
                    "file_type": get_document_type(
                        prepared_file
                    ),
                    "text": "",
                    "success": True,
                    "error": None,
                    "archive_path": (
                        str(downloaded_file)
                        if is_archive_url(url)
                        else ""
                    ),
                })

        except Exception as error:
            results.append({
                "url": url,
                "local_path": "",
                "normalized_path": "",
                "file_type": "",
                "text": "",
                "success": False,
                "error": str(error),
                "archive_path": "",
            })

    return results
